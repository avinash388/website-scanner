"""
Website Text Scanner  v2.0
Run: python app.py

Search terms come from:
  • Excel / CSV file upload  (local file)
  • SharePoint List  or  Excel file stored on SharePoint
"""

import sys
import io
import time
import threading
from collections import defaultdict
from datetime import datetime
from urllib.parse import urljoin, urlparse, urldefrag
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, scrolledtext

# ── Auto-install core deps ─────────────────────────────────────────────────────
try:
    import requests
    from bs4 import BeautifulSoup
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install",
                           "requests", "beautifulsoup4"])
    import requests
    from bs4 import BeautifulSoup

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

# ── Optional SharePoint client (installed on demand) ──────────────────────────
_SP_AVAILABLE = False
try:
    from office365.sharepoint.client_context import ClientContext
    from office365.runtime.auth.user_credential import UserCredential
    from office365.runtime.auth.client_credential import ClientCredential
    _SP_AVAILABLE = True
except ImportError:
    pass


def _ensure_sp():
    """Install Office365-REST-Python-Client if absent. Returns True on success."""
    global _SP_AVAILABLE
    if _SP_AVAILABLE:
        return True
    import subprocess
    try:
        subprocess.run(
            [sys.executable, "-m", "pip", "install",
             "Office365-REST-Python-Client"],
            check=True, capture_output=True, timeout=180)
        from office365.sharepoint.client_context import ClientContext      # noqa
        from office365.runtime.auth.user_credential import UserCredential  # noqa
        from office365.runtime.auth.client_credential import ClientCredential  # noqa
        _SP_AVAILABLE = True
        return True
    except Exception:
        return False


# ── Crawler ────────────────────────────────────────────────────────────────────

class WebsiteScanner:
    def __init__(self, base_url, search_texts, delay, max_pages,
                 on_progress, on_done):
        self.base_url     = base_url.rstrip("/")
        self.base_domain  = urlparse(base_url).netloc
        self.search_texts = search_texts
        self.delay        = delay
        self.max_pages    = max_pages
        self.on_progress  = on_progress
        self.on_done      = on_done
        self._stop        = False
        self.visited      = set()
        self.to_visit     = [self.base_url]
        self.results      = []
        self.errors       = []
        self.session      = requests.Session()
        self.session.headers.update({"User-Agent": "WebsiteScanner/2.0"})

    def stop(self): self._stop = True

    def _is_internal(self, url):
        d = urlparse(url).netloc
        return d == "" or d == self.base_domain

    def _normalize(self, url):
        url, _ = urldefrag(url)
        return url.rstrip("/")

    def _extract_links(self, soup, page_url):
        links = set()
        for tag in soup.find_all("a", href=True):
            href = tag["href"].strip()
            if any(href.startswith(p)
                   for p in ("mailto:", "tel:", "javascript:", "#")):
                continue
            norm = self._normalize(urljoin(page_url, href))
            if self._is_internal(norm) and norm.startswith("http"):
                links.add(norm)
        return links

    def run(self):
        while (self.to_visit and
               len(self.visited) < self.max_pages and
               not self._stop):
            url = self.to_visit.pop(0)
            if url in self.visited:
                continue
            self.visited.add(url)
            self.on_progress(f"[{len(self.visited)}] Checking: {url}")
            try:
                resp = self.session.get(url, timeout=10)
                if "text/html" not in resp.headers.get("Content-Type", ""):
                    continue
                resp.raise_for_status()
            except requests.RequestException as e:
                self.errors.append({"url": url, "error": str(e)})
                self.on_progress(f"     ⚠  Error: {e}")
                continue
            soup = BeautifulSoup(resp.text, "html.parser")
            title_tag = soup.find("title")
            title   = title_tag.get_text(strip=True) if title_tag else "(no title)"
            body    = soup.get_text(" ", strip=True).lower()
            matches = [t for t in self.search_texts if t.lower() in body]
            if matches:
                self.results.append({"url": url, "title": title,
                                     "matches": matches})
                self.on_progress(f"     ✔  MATCH — {', '.join(matches)}")
            for link in self._extract_links(soup, url):
                if link not in self.visited and link not in self.to_visit:
                    self.to_visit.append(link)
            time.sleep(self.delay)
        self.on_done(self.results, self.errors, len(self.visited))


# ── HTML report ────────────────────────────────────────────────────────────────

def build_html_report(base_url, search_texts, results, errors, pages_scanned):
    found_terms = defaultdict(list)
    for r in results:
        for m in r["matches"]:
            found_terms[m].append(r["url"])
    not_found = [t for t in search_texts if t not in found_terms]

    def _badges(matches):
        return "".join(f'<span class="badge">{m}</span>' for m in matches)

    rows = "".join(
        f'<tr><td class="num">{i}</td>'
        f'<td><a href="{r["url"]}" target="_blank">{r["url"]}</a>'
        f'<br><small>{r["title"]}</small></td>'
        f'<td>{_badges(r["matches"])}</td></tr>'
        for i, r in enumerate(results, 1))

    summary = "".join(
        f'<tr><td class="{"found" if t in found_terms else "notfound"}">'
        f'{"✔" if t in found_terms else "✘"} {t}</td>'
        f'<td>{"Found on "+str(len(found_terms[t]))+" page(s)" if t in found_terms else "Not found"}</td></tr>'
        for t in search_texts)

    err_html = ""
    if errors:
        items = "".join(f"<li>{e['url']} — {e['error']}</li>" for e in errors)
        err_html = (f"<details><summary>Errors ({len(errors)})</summary>"
                    f"<ul>{items}</ul></details>")

    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<title>Scan Report — {base_url}</title>
<style>
*{{box-sizing:border-box;margin:0;padding:0}}
body{{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;background:#f5f7fa;color:#222;padding:24px}}
h1{{font-size:1.5rem;margin-bottom:4px}}.meta{{color:#666;font-size:.85rem;margin-bottom:24px}}
.cards{{display:flex;gap:16px;flex-wrap:wrap;margin-bottom:28px}}
.card{{background:#fff;border-radius:8px;padding:16px 24px;box-shadow:0 1px 4px rgba(0,0,0,.1);min-width:130px}}
.card .n{{font-size:2rem;font-weight:700;color:#2563eb}}.card .l{{font-size:.8rem;color:#666;margin-top:2px}}
h2{{font-size:1.1rem;margin-bottom:12px;color:#333}}
table{{width:100%;border-collapse:collapse;background:#fff;border-radius:8px;overflow:hidden;box-shadow:0 1px 4px rgba(0,0,0,.1);margin-bottom:28px}}
th{{background:#1e40af;color:#fff;padding:10px 14px;text-align:left;font-size:.85rem}}
td{{padding:10px 14px;border-bottom:1px solid #eee;font-size:.875rem;vertical-align:top}}
tr:last-child td{{border-bottom:none}}.num{{color:#999;width:40px}}
a{{color:#2563eb;text-decoration:none}}a:hover{{text-decoration:underline}}small{{color:#888}}
.badge{{display:inline-block;background:#dbeafe;color:#1d4ed8;border-radius:4px;padding:2px 8px;font-size:.75rem;margin:2px 2px 0 0}}
.found{{color:#166534;font-weight:600;padding:6px 14px}}.notfound{{color:#991b1b;font-weight:600;padding:6px 14px}}
.none{{color:#888;font-style:italic;padding:16px}}
details{{background:#fff;border-radius:8px;padding:12px 16px;box-shadow:0 1px 4px rgba(0,0,0,.1);margin-bottom:24px}}
details summary{{cursor:pointer;font-weight:600;color:#c0392b}}
details ul{{margin-top:8px;padding-left:20px;font-size:.8rem;color:#555}}
</style></head><body>
<h1>Website Text Scan Report</h1>
<p class="meta">Base URL: <strong>{base_url}</strong> &nbsp;|&nbsp; Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
<div class="cards">
  <div class="card"><div class="n">{pages_scanned}</div><div class="l">Pages scanned</div></div>
  <div class="card"><div class="n">{len(results)}</div><div class="l">Pages with matches</div></div>
  <div class="card"><div class="n">{len(search_texts)}</div><div class="l">Search terms</div></div>
  <div class="card"><div class="n">{len(not_found)}</div><div class="l">Terms not found</div></div>
</div>
<h2>Search Term Summary</h2>
<table><thead><tr><th>Term</th><th>Result</th></tr></thead><tbody>{summary}</tbody></table>
<h2>Matched Pages ({len(results)})</h2>
<table><thead><tr><th>#</th><th>URL / Page Title</th><th>Matched Terms</th></tr></thead>
<tbody>{rows if results else '<tr><td colspan="3" class="none">No pages matched.</td></tr>'}</tbody></table>
{err_html}</body></html>"""


# ── UI constants ───────────────────────────────────────────────────────────────

ACCENT  = "#2563eb"
BG      = "#f0f4f8"
CARD_BG = "#ffffff"
FG      = "#1e293b"
MUTED   = "#64748b"
SUCCESS = "#16a34a"
DANGER  = "#dc2626"
SP_BLUE = "#0078d4"
FONT    = ("Segoe UI", 10)
FONT_B  = ("Segoe UI", 10, "bold")
FONT_S  = ("Segoe UI", 9)


# ── App ────────────────────────────────────────────────────────────────────────

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Website Text Scanner")
        self.configure(bg=BG)
        self.resizable(True, True)
        self.minsize(700, 700)

        # scan state
        self._scanner       = None
        self._thread        = None
        self._scan_results  = None
        self._scan_errors   = None
        self._pages_scanned = 0

        # unified search-terms list (all sources feed here)
        self._all_terms: list = []

        # excel state
        self._excel_wb     = None
        self._excel_path   = None
        self._excel_is_csv = False

        self._build_ui()
        self.update_idletasks()
        w, h = 760, 820
        x = (self.winfo_screenwidth()  - w) // 2
        y = (self.winfo_screenheight() - h) // 2
        self.geometry(f"{w}x{h}+{x}+{y}")

    # ══════════════════════════════════════════════════════════════════════════
    # Layout
    # ══════════════════════════════════════════════════════════════════════════

    def _build_ui(self):
        # Header
        hdr = tk.Frame(self, bg=ACCENT)
        hdr.pack(fill="x")
        tk.Label(hdr, text="  🔍 Website Text Scanner",
                 bg=ACCENT, fg="white",
                 font=("Segoe UI", 14, "bold"), pady=11).pack(side="left")

        # Scrollable body
        outer  = tk.Frame(self, bg=BG)
        outer.pack(fill="both", expand=True)
        canvas = tk.Canvas(outer, bg=BG, highlightthickness=0)
        vsb    = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        self._body = tk.Frame(canvas, bg=BG)
        cw = canvas.create_window((0, 0), window=self._body, anchor="nw")

        self._body.bind("<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>",
            lambda e: canvas.itemconfig(cw, width=e.width))
        canvas.bind_all("<MouseWheel>",
            lambda e: canvas.yview_scroll(int(-1*(e.delta/120)), "units"))

        main = tk.Frame(self._body, bg=BG)
        main.pack(fill="both", expand=True, padx=20, pady=16)

        self._build_url_section(main)
        self._build_terms_section(main)
        self._build_options_section(main)
        self._build_buttons(main)
        self._build_progress(main)
        self._build_log(main)

    def _section(self, parent, title):
        tk.Label(parent, text=title, bg=BG, fg=FG,
                 font=FONT_B, anchor="w").pack(fill="x", pady=(8, 4))

    def _card(self, parent):
        return tk.Frame(parent, bg=CARD_BG, relief="flat",
                        highlightbackground="#cbd5e1", highlightthickness=1)

    # ── Section 1: URL ────────────────────────────────────────────────────────

    def _build_url_section(self, parent):
        self._section(parent, "1  Website URL")
        card = self._card(parent)
        card.pack(fill="x", pady=(0, 12))
        self._url_var = tk.StringVar()
        e = tk.Entry(card, textvariable=self._url_var,
                     font=FONT, bd=0, fg=FG, bg=CARD_BG, insertbackground=FG)
        e.pack(fill="x", padx=12, pady=10)
        e.insert(0, "https://")

    # ══════════════════════════════════════════════════════════════════════════
    # Section 2: Search Terms  (tabbed: Excel Upload | SharePoint)
    # ══════════════════════════════════════════════════════════════════════════

    def _build_terms_section(self, parent):
        self._section(parent, "2  Search Terms")

        # ── tab notebook ──────────────────────────────────────────────────────
        nb_card = self._card(parent)
        nb_card.pack(fill="x", pady=(0, 6))

        sty = ttk.Style()
        sty.configure("T.TNotebook",
                      background=CARD_BG, borderwidth=0, tabmargins=[0, 0, 0, 0])
        sty.configure("T.TNotebook.Tab", padding=[16, 8], font=FONT_S)
        sty.map("T.TNotebook.Tab",
                background=[("selected", ACCENT), ("!selected", "#e2e8f0")],
                foreground=[("selected", "white"),  ("!selected", FG)])

        nb = ttk.Notebook(nb_card, style="T.TNotebook")
        nb.pack(fill="x", padx=2, pady=2)

        t1 = tk.Frame(nb, bg=CARD_BG)
        nb.add(t1, text="  📊  Excel Upload  ")
        self._build_excel_tab(t1)

        t2 = tk.Frame(nb, bg=CARD_BG)
        nb.add(t2, text="  ☁  SharePoint  ")
        self._build_sharepoint_tab(t2)

        # ── Unified loaded-terms list ─────────────────────────────────────────
        hdr = tk.Frame(parent, bg=BG)
        hdr.pack(fill="x", pady=(4, 2))
        self._terms_count_var = tk.StringVar(value="Loaded Search Terms (0)")
        tk.Label(hdr, textvariable=self._terms_count_var,
                 bg=BG, fg=FG, font=FONT_B, anchor="w").pack(side="left")
        tk.Button(hdr, text="Clear All",
                  bg="#fef2f2", fg=DANGER, bd=0, cursor="hand2",
                  font=FONT_S, padx=8, pady=3,
                  activebackground="#fee2e2", activeforeground=DANGER,
                  command=self._clear_terms).pack(side="right")

        list_card = self._card(parent)
        list_card.pack(fill="x", pady=(0, 12))

        lb_frame = tk.Frame(list_card, bg=CARD_BG)
        lb_frame.pack(fill="x")

        self._terms_lb = tk.Listbox(
            lb_frame, font=FONT, bg=CARD_BG, fg=FG, bd=0,
            selectbackground="#dbeafe", selectforeground=FG,
            highlightthickness=0, height=5, relief="flat",
            activestyle="none")
        vsb2 = ttk.Scrollbar(lb_frame, orient="vertical",
                              command=self._terms_lb.yview)
        self._terms_lb.configure(yscrollcommand=vsb2.set)
        self._terms_lb.pack(side="left", fill="both", expand=True,
                            padx=(8, 0), pady=4)
        vsb2.pack(side="right", fill="y", pady=4, padx=(0, 4))

        self._terms_lb.bind("<Delete>",    lambda e: self._remove_selected())
        self._terms_lb.bind("<BackSpace>", lambda e: self._remove_selected())

        ctx_menu = tk.Menu(self._terms_lb, tearoff=0)
        ctx_menu.add_command(label="Remove Selected", command=self._remove_selected)
        ctx_menu.add_command(label="Clear All",       command=self._clear_terms)

        def _show_ctx(e):
            self._terms_lb.selection_clear(0, "end")
            self._terms_lb.selection_set(self._terms_lb.nearest(e.y))
            try:    ctx_menu.tk_popup(e.x_root, e.y_root)
            finally: ctx_menu.grab_release()

        self._terms_lb.bind("<Button-3>", _show_ctx)

        tk.Label(list_card,
                 text="Select a term and press Delete to remove  |  right-click for menu",
                 bg=CARD_BG, fg="#94a3b8",
                 font=("Segoe UI", 8), anchor="w").pack(
            fill="x", padx=8, pady=(0, 4))

    # ── Excel Upload tab ──────────────────────────────────────────────────────

    def _build_excel_tab(self, parent):
        inner = tk.Frame(parent, bg=CARD_BG)
        inner.pack(fill="x", padx=14, pady=12)

        # File picker row
        r1 = tk.Frame(inner, bg=CARD_BG)
        r1.pack(fill="x", pady=(0, 8))
        tk.Label(r1, text="File:", bg=CARD_BG, fg=FG,
                 font=FONT, width=9, anchor="w").pack(side="left")
        self._xl_file_var = tk.StringVar(value="No file selected")
        tk.Label(r1, textvariable=self._xl_file_var,
                 bg=BG, fg=MUTED, font=FONT_S, anchor="w",
                 width=36).pack(side="left", fill="x", expand=True, padx=(4, 8))
        tk.Button(r1, text="Browse…",
                  bg=ACCENT, fg="white", bd=0, cursor="hand2",
                  font=FONT_S, padx=12, pady=5,
                  activebackground="#1d4ed8", activeforeground="white",
                  command=self._xl_browse).pack(side="right")

        # Sheet + Column selectors
        r2 = tk.Frame(inner, bg=CARD_BG)
        r2.pack(fill="x", pady=(0, 8))
        tk.Label(r2, text="Sheet:", bg=CARD_BG, fg=FG,
                 font=FONT, width=9, anchor="w").pack(side="left")
        self._xl_sheet_var = tk.StringVar()
        self._xl_sheet_cb  = ttk.Combobox(r2, textvariable=self._xl_sheet_var,
                                          font=FONT_S, state="readonly", width=17)
        self._xl_sheet_cb.pack(side="left", padx=(4, 16))
        self._xl_sheet_cb.bind("<<ComboboxSelected>>", self._xl_sheet_changed)

        tk.Label(r2, text="Column:", bg=CARD_BG, fg=FG,
                 font=FONT).pack(side="left")
        self._xl_col_var = tk.StringVar()
        self._xl_col_cb  = ttk.Combobox(r2, textvariable=self._xl_col_var,
                                        font=FONT_S, state="readonly", width=17)
        self._xl_col_cb.pack(side="left", padx=(4, 0))

        # Info
        self._xl_info_var = tk.StringVar(
            value="Browse an Excel (.xlsx) or CSV file — terms are loaded from the selected column")
        tk.Label(inner, textvariable=self._xl_info_var,
                 bg=CARD_BG, fg=MUTED, font=FONT_S,
                 anchor="w", wraplength=600).pack(fill="x", pady=(0, 8))

        tk.Button(inner, text="📥  Add Terms from File",
                  bg=SUCCESS, fg="white", bd=0, cursor="hand2",
                  font=FONT_S, padx=14, pady=6,
                  activebackground="#15803d", activeforeground="white",
                  command=self._xl_load).pack(anchor="w")

    def _xl_browse(self):
        path = filedialog.askopenfilename(
            title="Select Excel or CSV",
            filetypes=[("Excel / CSV", "*.xlsx *.xls *.csv"),
                       ("All Files", "*.*")])
        if not path:
            return
        self._excel_path = path
        fname = path.replace("\\", "/").split("/")[-1]
        self._xl_file_var.set(fname)
        try:
            if path.lower().endswith(".csv"):
                self._excel_is_csv = True
                self._excel_wb     = None
                import csv
                with open(path, newline="", encoding="utf-8-sig") as f:
                    headers = next(csv.reader(f), [])
                headers = headers or ["Column 1"]
                self._xl_sheet_cb.configure(values=["(CSV)"])
                self._xl_sheet_var.set("(CSV)")
                self._xl_col_cb.configure(values=headers)
                self._xl_col_var.set(headers[0])
                self._xl_info_var.set(
                    f"CSV loaded — {len(headers)} column(s). "
                    "Choose the column that contains search terms.")
            else:
                self._excel_is_csv = False
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                self._excel_wb = wb
                sheets = wb.sheetnames
                self._xl_sheet_cb.configure(values=sheets)
                self._xl_sheet_var.set(sheets[0])
                self._xl_sheet_changed()
        except Exception as ex:
            messagebox.showerror("Load Error", f"Could not open file:\n{ex}")

    def _xl_sheet_changed(self, _=None):
        if not self._excel_wb:
            return
        ws = self._excel_wb[self._xl_sheet_var.get()]
        headers = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            for c in row:
                headers.append(str(c) if c is not None
                               else f"Col {len(headers)+1}")
        headers = headers or [f"Col {i+1}" for i in range(ws.max_column or 1)]
        self._xl_col_cb.configure(values=headers)
        self._xl_col_var.set(headers[0])
        self._xl_info_var.set(
            f"Sheet '{self._xl_sheet_var.get()}' — "
            f"{ws.max_row or 0} rows · {len(headers)} column(s). "
            "Choose the column that contains search terms.")

    def _xl_load(self):
        if not self._excel_path:
            messagebox.showwarning("No File", "Browse for a file first.")
            return
        col = self._xl_col_var.get()
        if not col:
            messagebox.showwarning("No Column", "Select a column first.")
            return
        try:
            terms = []
            if self._excel_is_csv:
                import csv
                with open(self._excel_path, newline="",
                          encoding="utf-8-sig") as f:
                    for row in csv.DictReader(f):
                        v = row.get(col, "").strip()
                        if v:
                            terms.append(v)
            else:
                ws = self._excel_wb[self._xl_sheet_var.get()]
                hdr = next(ws.iter_rows(min_row=1, max_row=1,
                                        values_only=True), [])
                try:
                    ci = list(hdr).index(col)
                except ValueError:
                    messagebox.showerror("Column Error",
                        f"Column '{col}' not found.")
                    return
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if ci < len(row) and row[ci] is not None:
                        v = str(row[ci]).strip()
                        if v:
                            terms.append(v)
            if not terms:
                messagebox.showinfo("Empty", "No terms found in that column.")
                return
            added = self._add_terms(terms)
            self._show_added(added, len(terms))
        except Exception as ex:
            messagebox.showerror("Read Error", f"Failed to read terms:\n{ex}")

    # ── SharePoint tab ────────────────────────────────────────────────────────

    def _build_sharepoint_tab(self, parent):
        inner = tk.Frame(parent, bg=CARD_BG)
        inner.pack(fill="x", padx=14, pady=12)

        def field_row(container, label_text, var, show="", lw=14):
            row = tk.Frame(container, bg=CARD_BG)
            row.pack(fill="x", pady=2)
            tk.Label(row, text=label_text, bg=CARD_BG, fg=FG,
                     font=FONT, width=lw, anchor="w").pack(side="left")
            tk.Entry(row, textvariable=var, font=FONT, bd=0,
                     fg=FG, bg=BG, insertbackground=FG,
                     relief="flat", show=show).pack(
                side="left", fill="x", expand=True, ipady=4, padx=(4, 0))

        # Site URL
        self._sp_url_var = tk.StringVar(
            value="https://yourcompany.sharepoint.com/sites/SiteName")
        field_row(inner, "Site URL:", self._sp_url_var)

        # Auth method
        auth_row = tk.Frame(inner, bg=CARD_BG)
        auth_row.pack(fill="x", pady=(8, 2))
        tk.Label(auth_row, text="Auth Method:", bg=CARD_BG, fg=FG,
                 font=FONT, width=14, anchor="w").pack(side="left")
        self._sp_auth_var = tk.StringVar(value="user")
        for val, lbl in [("user", "User Credentials"),
                         ("app",  "App (Client ID / Secret)")]:
            tk.Radiobutton(auth_row, text=lbl,
                           variable=self._sp_auth_var, value=val,
                           bg=CARD_BG, fg=FG, font=FONT_S,
                           activebackground=CARD_BG, selectcolor=CARD_BG,
                           command=self._sp_toggle_auth).pack(
                side="left", padx=(4, 14))

        # User-credentials frame
        self._sp_user_frame = tk.Frame(inner, bg=CARD_BG)
        self._sp_user_var = tk.StringVar()
        self._sp_pass_var = tk.StringVar()
        field_row(self._sp_user_frame, "Username:", self._sp_user_var)
        field_row(self._sp_user_frame, "Password:", self._sp_pass_var, show="•")

        # App-credentials frame
        self._sp_app_frame  = tk.Frame(inner, bg=CARD_BG)
        self._sp_cid_var    = tk.StringVar()
        self._sp_sec_var    = tk.StringVar()
        self._sp_tenant_var = tk.StringVar()
        field_row(self._sp_app_frame, "Client ID:",     self._sp_cid_var)
        field_row(self._sp_app_frame, "Client Secret:", self._sp_sec_var, show="•")
        field_row(self._sp_app_frame, "Tenant ID:",     self._sp_tenant_var)

        # Source type
        src_row = tk.Frame(inner, bg=CARD_BG)
        src_row.pack(fill="x", pady=(8, 2))
        tk.Label(src_row, text="Source:", bg=CARD_BG, fg=FG,
                 font=FONT, width=14, anchor="w").pack(side="left")
        self._sp_src_var = tk.StringVar(value="list")
        for val, lbl in [("list",  "SharePoint List"),
                         ("excel", "Excel File on SharePoint")]:
            tk.Radiobutton(src_row, text=lbl,
                           variable=self._sp_src_var, value=val,
                           bg=CARD_BG, fg=FG, font=FONT_S,
                           activebackground=CARD_BG, selectcolor=CARD_BG,
                           command=self._sp_toggle_src).pack(
                side="left", padx=(4, 14))

        # List fields
        self._sp_list_frame   = tk.Frame(inner, bg=CARD_BG)
        self._sp_list_var     = tk.StringVar(value="Search Terms")
        self._sp_list_col_var = tk.StringVar(value="Title")
        field_row(self._sp_list_frame, "List Name:",   self._sp_list_var)
        field_row(self._sp_list_frame, "Column Name:", self._sp_list_col_var)

        # Excel-on-SP fields
        self._sp_xl_frame    = tk.Frame(inner, bg=CARD_BG)
        self._sp_xl_path_var = tk.StringVar(
            value="/sites/SiteName/Shared Documents/terms.xlsx")
        self._sp_xl_col_var  = tk.StringVar(value="Term")
        field_row(self._sp_xl_frame, "File Path:",    self._sp_xl_path_var)
        field_row(self._sp_xl_frame, "Column Name:", self._sp_xl_col_var)

        # Fetch button + status
        btn_row = tk.Frame(inner, bg=CARD_BG)
        btn_row.pack(fill="x", pady=(10, 0))
        self._sp_btn = tk.Button(
            btn_row, text="☁  Fetch Terms from SharePoint",
            bg=SP_BLUE, fg="white", bd=0, cursor="hand2",
            font=FONT_S, padx=14, pady=7,
            activebackground="#006cbd", activeforeground="white",
            command=self._sp_start)
        self._sp_btn.pack(side="left")
        self._sp_status_var = tk.StringVar(value="Ready")
        self._sp_lbl = tk.Label(btn_row, textvariable=self._sp_status_var,
                                bg=CARD_BG, fg=MUTED, font=FONT_S)
        self._sp_lbl.pack(side="left", padx=(12, 0))

        # set initial visibility
        self._sp_toggle_auth()
        self._sp_toggle_src()

    def _sp_toggle_auth(self):
        self._sp_app_frame.pack_forget()
        self._sp_user_frame.pack_forget()
        if self._sp_auth_var.get() == "user":
            self._sp_user_frame.pack(fill="x", pady=(0, 4))
        else:
            self._sp_app_frame.pack(fill="x", pady=(0, 4))

    def _sp_toggle_src(self):
        self._sp_xl_frame.pack_forget()
        self._sp_list_frame.pack_forget()
        if self._sp_src_var.get() == "list":
            self._sp_list_frame.pack(fill="x", pady=(0, 4))
        else:
            self._sp_xl_frame.pack(fill="x", pady=(0, 4))

    def _sp_start(self):
        site = self._sp_url_var.get().strip()
        if not site or "yourcompany" in site:
            messagebox.showwarning("Missing URL",
                "Enter your SharePoint site URL first.")
            return
        if not _SP_AVAILABLE:
            self._sp_status_var.set("Installing library…")
            self._sp_lbl.config(fg="#f59e0b")
            self.update_idletasks()
            if not _ensure_sp():
                messagebox.showerror(
                    "Package Required",
                    "Could not auto-install Office365-REST-Python-Client.\n\n"
                    "Install manually and restart:\n"
                    "  pip install Office365-REST-Python-Client")
                self._sp_status_var.set("Library missing — see error")
                self._sp_lbl.config(fg=DANGER)
                return
        self._sp_btn.config(state="disabled")
        self._sp_status_var.set("Connecting…")
        self._sp_lbl.config(fg="#f59e0b")
        threading.Thread(target=self._sp_thread, daemon=True).start()

    def _sp_thread(self):
        try:
            from office365.sharepoint.client_context import ClientContext
            from office365.runtime.auth.user_credential import UserCredential
            from office365.runtime.auth.client_credential import ClientCredential

            site = self._sp_url_var.get().strip()
            if self._sp_auth_var.get() == "user":
                cred = UserCredential(self._sp_user_var.get().strip(),
                                      self._sp_pass_var.get())
            else:
                cred = ClientCredential(self._sp_cid_var.get().strip(),
                                        self._sp_sec_var.get().strip())
            ctx = ClientContext(site).with_credentials(cred)

            self.after(0, lambda: self._sp_status_var.set("Fetching…"))

            terms = (self._sp_fetch_list(ctx)
                     if self._sp_src_var.get() == "list"
                     else self._sp_fetch_excel(ctx))

            added = self._add_terms(terms)

            def _ok():
                self._sp_btn.config(state="normal")
                self._sp_status_var.set(
                    f"✔ {added} new term(s) added "
                    f"({len(terms)-added} duplicate(s) skipped)")
                self._sp_lbl.config(fg=SUCCESS)
                messagebox.showinfo("SharePoint",
                    f"Fetched {len(terms)} term(s) from SharePoint.\n"
                    f"{added} new, {len(terms)-added} duplicate(s) skipped.")
            self.after(0, _ok)

        except Exception as ex:
            err = str(ex)
            def _fail():
                self._sp_btn.config(state="normal")
                self._sp_status_var.set(f"Error: {err[:55]}")
                self._sp_lbl.config(fg=DANGER)
                messagebox.showerror("SharePoint Error",
                    f"Could not fetch terms:\n\n{err}")
            self.after(0, _fail)

    def _sp_fetch_list(self, ctx):
        name = self._sp_list_var.get().strip()
        col  = self._sp_list_col_var.get().strip()
        if not name or not col:
            raise ValueError("List name and column name are required.")
        self.after(0, lambda: self._sp_status_var.set(
            f"Reading list '{name}'…"))
        items = (ctx.web.lists
                    .get_by_title(name)
                    .items.select([col])
                    .top(5000)
                    .get()
                    .execute_query())
        return [str(it.properties[col]).strip()
                for it in items
                if it.properties.get(col)
                and str(it.properties[col]).strip()]

    def _sp_fetch_excel(self, ctx):
        path = self._sp_xl_path_var.get().strip()
        col  = self._sp_xl_col_var.get().strip()
        if not path or not col:
            raise ValueError("File path and column name are required.")
        self.after(0, lambda: self._sp_status_var.set(
            "Downloading file from SharePoint…"))
        content = (ctx.web
                      .get_file_by_server_relative_url(path)
                      .read()
                      .execute_query())
        wb = openpyxl.load_workbook(io.BytesIO(content.value),
                                    read_only=True, data_only=True)
        ws = wb.active
        hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
        try:
            ci = list(hdr).index(col)
        except ValueError:
            raise ValueError(
                f"Column '{col}' not found. Available: {[h for h in hdr if h]}")
        return [str(row[ci]).strip()
                for row in ws.iter_rows(min_row=2, values_only=True)
                if ci < len(row) and row[ci] is not None
                and str(row[ci]).strip()]

    # ══════════════════════════════════════════════════════════════════════════
    # Unified term management
    # ══════════════════════════════════════════════════════════════════════════

    def _add_terms(self, terms):
        added = 0
        for t in terms:
            t = str(t).strip()
            if t and t not in self._all_terms:
                self._all_terms.append(t)
                added += 1
        self._refresh_lb()
        return added

    def _remove_selected(self):
        sel = self._terms_lb.curselection()
        if sel:
            self._all_terms.pop(sel[0])
            self._refresh_lb()

    def _clear_terms(self):
        if self._all_terms and messagebox.askyesno(
                "Clear All", "Remove all loaded search terms?"):
            self._all_terms.clear()
            self._refresh_lb()

    def _refresh_lb(self):
        self._terms_lb.delete(0, "end")
        for t in self._all_terms:
            self._terms_lb.insert("end", f"  {t}")
        self._terms_count_var.set(
            f"Loaded Search Terms ({len(self._all_terms)})")

    def _get_terms(self):
        return list(self._all_terms)

    def _show_added(self, added, total):
        skipped = total - added
        if added:
            msg = f"✔ {added} term(s) added to the list."
            if skipped:
                msg += f"  ({skipped} duplicate(s) skipped)"
            messagebox.showinfo("Terms Added", msg)
        else:
            messagebox.showinfo("No New Terms",
                "All terms are already in the list (no duplicates added).")

    # ══════════════════════════════════════════════════════════════════════════
    # Section 3: Options
    # ══════════════════════════════════════════════════════════════════════════

    def _build_options_section(self, parent):
        self._section(parent, "3  Options")
        card = self._card(parent)
        card.pack(fill="x", pady=(0, 14))
        inner = tk.Frame(card, bg=CARD_BG)
        inner.pack(fill="x", padx=12, pady=8)

        tk.Label(inner, text="Max pages:", bg=CARD_BG, fg=FG,
                 font=FONT).grid(row=0, column=0, sticky="w", pady=4)
        self._max_pages_var = tk.StringVar(value="200")
        tk.Entry(inner, textvariable=self._max_pages_var, width=8,
                 font=FONT, bg=BG, relief="flat").grid(
            row=0, column=1, sticky="w", padx=8)

        tk.Label(inner, text="Delay (sec):", bg=CARD_BG, fg=FG,
                 font=FONT).grid(row=0, column=2, sticky="w", padx=(24, 0))
        self._delay_var = tk.StringVar(value="0.3")
        tk.Entry(inner, textvariable=self._delay_var, width=6,
                 font=FONT, bg=BG, relief="flat").grid(
            row=0, column=3, sticky="w", padx=8)

    # ══════════════════════════════════════════════════════════════════════════
    # Buttons / progress / log
    # ══════════════════════════════════════════════════════════════════════════

    def _build_buttons(self, parent):
        row = tk.Frame(parent, bg=BG)
        row.pack(fill="x", pady=(0, 12))

        self._start_btn = tk.Button(
            row, text="  ▶  Start Scan",
            bg=ACCENT, fg="white", font=FONT_B, bd=0, cursor="hand2",
            padx=20, pady=8,
            activebackground="#1d4ed8", activeforeground="white",
            command=self._start_scan)
        self._start_btn.pack(side="left")

        self._stop_btn = tk.Button(
            row, text="  ■  Stop",
            bg="#e2e8f0", fg=MUTED, font=FONT_B, bd=0, cursor="hand2",
            padx=20, pady=8, state="disabled", command=self._stop_scan)
        self._stop_btn.pack(side="left", padx=8)

        self._save_btn = tk.Button(
            row, text="  ↓  Save Report",
            bg="#e2e8f0", fg=MUTED, font=FONT_B, bd=0, cursor="hand2",
            padx=20, pady=8, state="disabled", command=self._save_report)
        self._save_btn.pack(side="right")

    def _build_progress(self, parent):
        self._prog_var = tk.StringVar(value="Ready")
        tk.Label(parent, textvariable=self._prog_var,
                 bg=BG, fg=MUTED, font=FONT_S).pack(anchor="w")
        self._pbar = ttk.Progressbar(parent, mode="indeterminate", length=200)
        self._pbar.pack(fill="x", pady=(2, 8))

    def _build_log(self, parent):
        self._section(parent, "Log")
        self._log = scrolledtext.ScrolledText(
            parent, height=10, font=("Consolas", 9),
            bg="#0f172a", fg="#94a3b8", insertbackground="white",
            relief="flat", bd=0, state="disabled")
        self._log.pack(fill="both", expand=True)
        self._log.tag_config("match", foreground="#4ade80")
        self._log.tag_config("error", foreground="#f87171")
        self._log.tag_config("info",  foreground="#60a5fa")
        self._log.tag_config("done",  foreground="#facc15")

    def _log_write(self, msg, tag=""):
        def _do():
            self._log.config(state="normal")
            ts = datetime.now().strftime("%H:%M:%S")
            self._log.insert("end", f"[{ts}] {msg}\n", tag)
            self._log.see("end")
            self._log.config(state="disabled")
        self.after(0, _do)

    # ══════════════════════════════════════════════════════════════════════════
    # Scan control
    # ══════════════════════════════════════════════════════════════════════════

    def _start_scan(self):
        url = self._url_var.get().strip()
        if not url or url == "https://":
            messagebox.showwarning("Missing URL", "Enter a website URL.")
            return

        terms = self._get_terms()
        if not terms:
            messagebox.showwarning("No Terms",
                "Load at least one search term (Excel or SharePoint) "
                "before starting the scan.")
            return

        if not url.startswith("http"):
            url = "https://" + url

        try:
            max_pages = int(self._max_pages_var.get())
            delay     = float(self._delay_var.get())
        except ValueError:
            messagebox.showerror("Invalid Options",
                "Max pages must be an integer; delay must be a number.")
            return

        self._scan_results = None
        self._log.config(state="normal")
        self._log.delete("1.0", "end")
        self._log.config(state="disabled")
        self._save_btn.config(state="disabled", bg="#e2e8f0", fg=MUTED)
        self._start_btn.config(state="disabled")
        self._stop_btn.config(state="normal", bg="#fef2f2", fg=DANGER)
        self._pbar.start(12)
        self._prog_var.set("Scanning…")

        self._log_write(f"Target: {url}", "info")
        self._log_write(f"Terms ({len(terms)}): {', '.join(terms)}", "info")
        self._log_write(f"Max pages: {max_pages}  |  Delay: {delay}s", "info")
        self._log_write("─" * 60)

        self._scanner = WebsiteScanner(
            base_url=url, search_texts=terms,
            delay=delay, max_pages=max_pages,
            on_progress=self._on_progress, on_done=self._on_done)
        self._thread = threading.Thread(
            target=self._scanner.run, daemon=True)
        self._thread.start()

    def _stop_scan(self):
        if self._scanner:
            self._scanner.stop()
        self._log_write("Scan stopped by user.", "error")

    def _on_progress(self, msg):
        tag = "match" if "✔" in msg else ("error" if "⚠" in msg else "")
        self._log_write(msg, tag)
        self.after(0, lambda: self._prog_var.set(msg[:80]))

    def _on_done(self, results, errors, total):
        self._scan_results  = results
        self._scan_errors   = errors
        self._pages_scanned = total

        def _finish():
            self._pbar.stop()
            self._start_btn.config(state="normal")
            self._stop_btn.config(state="disabled", bg="#e2e8f0", fg=MUTED)
            self._prog_var.set(
                f"Done — {total} pages scanned, {len(results)} matched")
            self._log_write("─" * 60, "done")
            self._log_write(
                f"Scan complete: {total} pages · {len(results)} matched",
                "done")
            found = defaultdict(int)
            for r in results:
                for m in r["matches"]:
                    found[m] += 1
            for t in self._get_terms():
                if t in found:
                    self._log_write(
                        f"  ✔  '{t}'  → {found[t]} page(s)", "match")
                else:
                    self._log_write(f"  ✘  '{t}'  → NOT FOUND", "error")
            if results:
                self._save_btn.config(state="normal", bg=SUCCESS, fg="white")

        self.after(0, _finish)

    # ══════════════════════════════════════════════════════════════════════════
    # Save report
    # ══════════════════════════════════════════════════════════════════════════

    def _save_report(self):
        if not self._scan_results and self._scan_results is not None:
            messagebox.showinfo("No Matches",
                "No pages matched — nothing to save.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".html",
            filetypes=[("HTML Report", "*.html"), ("All Files", "*.*")],
            initialfile=(f"scan_report_"
                         f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"),
            title="Save Report")
        if not path:
            return
        html = build_html_report(
            self._url_var.get().strip(), self._get_terms(),
            self._scan_results, self._scan_errors, self._pages_scanned)
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
        self._log_write(f"Report saved: {path}", "info")
        messagebox.showinfo("Saved", f"Report saved to:\n{path}")


# ── Entry point ────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    app = App()
    app.mainloop()
